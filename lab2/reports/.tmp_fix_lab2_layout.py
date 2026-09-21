from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SRC = Path("/Users/ehoi/Library/Vivado/Workspace/ece2/lab2/reports/LAB2_실험전보고서.docx")
OUT = Path("/Users/ehoi/Library/Vivado/Workspace/ece2/lab2/reports/LAB2_실험전보고서_레이아웃수정본.docx")


def set_font(run, name="Nanum Gothic", size=None, color="000000"):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_style_font(style, name, size, bold=None, color="000000"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def set_cell_width(cell, width_inches):
    width_twips = str(round(width_inches * 1440))
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), width_twips)
    tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_width(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(round(total * 1440)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is not None:
        tbl_pr.remove(tbl_ind)

    grid = table._tbl.tblGrid
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        for col, width in zip(cols, widths):
            col.set(qn("w:w"), str(round(width * 1440)))

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def set_table_cell_margins(table, top=45, start=70, bottom=45, end=70):
    tbl_pr = table._tbl.tblPr
    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    for height in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(height)
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat_header:
        tbl_header = tr_pr.find(qn("w:tblHeader"))
        if tbl_header is None:
            tbl_header = OxmlElement("w:tblHeader")
            tr_pr.append(tbl_header)
        tbl_header.set(qn("w:val"), "true")


def insert_column_break_before(table):
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    p_style = OxmlElement("w:pStyle")
    p_style.set(qn("w:val"), "Normal")
    p_pr.append(p_style)
    p.append(p_pr)
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "column")
    run.append(br)
    p.append(run)
    table._tbl.addprevious(p)


doc = Document(SRC)

# Match the Lab 1 report's A4 page geometry while preserving the cover/body split.
for idx, section in enumerate(doc.sections):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "1" if idx == 0 else "2")
    cols.set(qn("w:space"), "425")

# Typography and paragraph rhythm copied from the Lab 1 visual system.
set_style_font(doc.styles["Normal"], "Nanum Gothic", 10.0, None)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.line_spacing = 1.12

set_style_font(doc.styles["Title"], "Nanum Gothic", 25.0, True)
doc.styles["Title"].paragraph_format.space_before = Pt(90)
doc.styles["Title"].paragraph_format.space_after = Pt(10)
doc.styles["Title"].paragraph_format.line_spacing = 1.0

set_style_font(doc.styles["Subtitle"], "Nanum Gothic", 12.0, False)
doc.styles["Subtitle"].paragraph_format.space_after = Pt(16)
doc.styles["Subtitle"].paragraph_format.line_spacing = 1.0

set_style_font(doc.styles["Heading 1"], "Nanum Gothic", 15.0, True)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.keep_together = True

set_style_font(doc.styles["Heading 2"], "Nanum Gothic", 11.5, True)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
doc.styles["Heading 2"].paragraph_format.keep_with_next = True
doc.styles["Heading 2"].paragraph_format.keep_together = True

if "Caption" in doc.styles:
    set_style_font(doc.styles["Caption"], "Nanum Gothic", 8.0, True, "333333")
    doc.styles["Caption"].paragraph_format.space_after = Pt(5)
    doc.styles["Caption"].paragraph_format.line_spacing = 1.0

# Compress only the five structural blank lines before the cover title.
for paragraph in doc.paragraphs[:5]:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)

for idx, paragraph in enumerate(doc.paragraphs):
    style_name = paragraph.style.name if paragraph.style else "Normal"
    has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
    text = paragraph.text.strip()

    if idx in (5, 6):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif 7 <= idx <= 10:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)

    if style_name == "Normal":
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.12
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            set_font(run, "Nanum Gothic", 10.0)

    if style_name in ("Heading 1", "Heading 2"):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.page_break_before = False
        for run in paragraph.runs:
            set_font(run, "Nanum Gothic", 15.0 if style_name == "Heading 1" else 11.5)
            run.bold = True

    if has_drawing:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True

    if text.startswith("그림 "):
        if "Caption" in doc.styles:
            paragraph.style = doc.styles["Caption"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
        for run in paragraph.runs:
            set_font(run, "Nanum Gothic", 8.0, "333333")
            run.bold = True

# Keep all figures safely inside one column and preserve their aspect ratios.
for shape in doc.inline_shapes:
    max_width = Inches(3.02)
    if shape.width > max_width:
        ratio = shape.height / shape.width
        shape.width = max_width
        shape.height = int(max_width * ratio)

# Fix all table grids and cell widths so Word cannot collapse the second column.
two_col_widths = {
    0: [1.06, 2.00],
    1: [0.86, 2.20],
    4: [0.78, 2.28],
    6: [0.78, 2.28],
    8: [0.78, 2.28],
    10: [0.78, 2.28],
    12: [0.78, 2.28],
    14: [0.78, 2.28],
    16: [0.78, 2.28],
    17: [0.92, 2.14],
}

code_tables = {2, 3, 5, 7, 9, 11, 13, 15}

for table_idx, table in enumerate(doc.tables):
    widths = two_col_widths.get(table_idx, [3.06])
    set_table_width(table, widths)
    set_table_borders(table)
    set_table_cell_margins(table)

    for row_idx, row in enumerate(table.rows):
        keep_row_together(row, repeat_header=(row_idx == 0 and table_idx not in code_tables))
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            is_header = row_idx == 0 and table_idx not in code_tables
            if table_idx in code_tables:
                set_cell_shading(cell, "F2F2F2")
            elif is_header:
                set_cell_shading(cell, "1F4E78")
            else:
                set_cell_shading(cell, "FFFFFF" if row_idx % 2 else "EAF2F8")

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_together = True
                if table_idx in code_tables:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif is_header or col_idx == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

                for run in paragraph.runs:
                    if table_idx in code_tables:
                        set_font(run, "Menlo", 7.2, "222222")
                    elif is_header:
                        set_font(run, "Nanum Gothic", 7.8, "FFFFFF")
                        run.bold = True
                    else:
                        set_font(run, "Nanum Gothic", 7.4, "222222")

# The two long 9-row summary tables previously split at the column boundary.
# Start each at the top of a fresh column so it remains visually contiguous.
insert_column_break_before(doc.tables[0])
insert_column_break_before(doc.tables[17])

# Normalize footer typography without changing its text or page fields.
for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_font(run, "Nanum Gothic", 7.0, "7F7F7F")

doc.save(OUT)
print(OUT)
