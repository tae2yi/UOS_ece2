from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/ehoi/Library/Vivado/Workspace/ece2/lab2')
OUT = ROOT / 'reports' / 'LAB2_실험전보고서.docx'
FONT_BODY = 'NanumBarunGothicOTF'
FONT_HEAD = 'NanumSquareOTF'
FONT_CODE = 'Menlo'
NAVY = '203B5A'
PALE = 'EAF1F7'
GRAY = 'D9D9D9'
ASSET_DIR = ROOT / 'reports' / 'evidence_crops'
COLUMN_WIDTH = 3.12


LABS = [
    {
        'id':'01','name':'4비트 업다운 카운터','module':'counter4','end':'356 ns','checks':'36',
        'goal':'동기식 리셋과 enable 제어를 갖는 4비트 업다운 카운터를 설계한다. down이 0이면 증가하고 1이면 감소하며, 4비트 범위 끝에서 자연스럽게 순환한다.',
        'structure':'상승 에지에서 rst를 가장 먼저 확인한다. rst가 1이면 value를 0으로 만들고, 그렇지 않을 때 enable이 1이면 down 방향에 따라 1을 더하거나 뺀다. enable이 0인 동안에는 값이 유지된다. 보드용 입력 전처리기와 top은 별도 RTL이며 이 TB는 counter4 코어만 검사한다.',
        'code':"""always @(posedge clk) begin
    if (rst) value <= 4'd0;
    else if (enable) begin
        if (down) value <= value - 4'd1;
        else value <= value + 4'd1;
    end
end""",
        'input':'TB의 clk 주기는 10 ns이다. reset 후 enable을 켜고 down=0으로 16번 검사해 0, 1, …, 15, 0을 확인한다. 이어 down=1로 16번 검사해 15, 14, …, 0, 15가 되는지 확인한다. enable을 내리면 값이 유지되어야 하며, reset과 enable이 함께 1일 때 reset이 우선해야 한다.',
        'normal':'LAB2_PASS counter4 checks=36. 종료 356 ns. 36개 자기검사에서 리셋, 양방향 증가·감소, 랩어라운드, hold 및 리셋 우선순위가 통과했다. input_frontend, XDC 핀 배치와 보드 동작을 검증한 결과는 아니다.',
        'normal_img':ROOT/'lab2_01_counter/evidence/normal/counter4_normal_waveform.png',
        'normal_caption':'전체 구간에서 15→0 및 0→15 순환과 종료 직전 hold·동기 리셋을 확인한다.',
        'mod':'교안에는 별도의 수정 실험이 제시되지 않는다. 정상 RTL과 TB를 그대로 두며, 임의의 고장을 주입하거나 수정 FAIL을 만들지 않았다.',
        'failure':'해당 없음','modified_img':None,
        'conclusion':'정상 파형은 각 상승 에지에서만 value가 변하고, enable이 꺼진 동안 값을 보존함을 보인다. 4비트 산술에 따라 증가·감소 경계에서 순환한다. 검증 범위는 코어 TB에 한정된다.',
        'pdf':'05.LAB2_01_COUNTER_VIVADO.pdf, 14–22쪽, 33–40쪽'
    },
    {
        'id':'02','name':'클록 분주기','module':'clock_divider','end':'436 ns','checks':'129',
        'goal':'입력 클록을 정수 비율로 분주해 관찰용 사각파 divided와 한 클록 폭 enable 신호 tick을 만든다.',
        'structure':'count는 DIVISOR−1까지 센 뒤 0으로 돌아간다. 짝수 DIVISOR에서 count가 절반 주기 경계에 도달하면 divided를 올리고 전체 주기 경계에서 내리므로 high·low 길이가 같다. tick은 count가 마지막 값일 때 1이며 다른 순차 회로의 enable로 사용한다. 테스트벤치에는 DIVISOR=10과 최소 허용값 2가 포함된다.',
        'code':"""localparam integer WIDTH = $clog2(DIVISOR);
reg [WIDTH-1:0] count;
assign tick = !rst && (count == DIVISOR - 1);
always @(posedge clk) begin
    if (rst) begin count <= 0; divided <= 1'b0; end
    else begin
        if (count == DIVISOR - 1) count <= 0;
        else count <= count + 1'b1;
        if (count == DIVISOR/2 - 1) divided <= 1'b1;
        else if (count == DIVISOR - 1) divided <= 1'b0;
    end
end""",
        'input':'TB clk는 10 ns 주기다. reset 뒤 DIVISOR=10에서 30주기 동안 divided의 위상·50% duty 및 마지막 count에서의 tick을 확인한다. DIVISOR=2 인스턴스에서도 매 2주기 출력과 enable pulse를 확인하고, 중간 reset 후 처음부터 다시 시작하는지 검사한다.',
        'normal':'LAB2_PASS clock_divider checks=129. 종료 436 ns. DIVISOR=10, 2의 duty와 위상, tick 시점·소비 pulse, reset 및 재시작 검사가 모두 통과했다.',
        'normal_img':ROOT/'lab2_02_clock_divider/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 VaporView 화면의 divided, tick 및 DIVISOR=2 출력·tick 파형.',
        'mod':'교안 지시에 따라 divided 상승 비교식을 count == DIVISOR/2 - 1에서 count == DIVISOR/2로 잠시 바꾸고 TB는 유지한다. DIVISOR=2에서는 상승 조건과 하강 조건이 겹쳐 첫 활성 에지 이후에도 div2가 낮게 남는다.',
        'failure':'16 ns · minimum divisor duty 검사. 기대 div2=1, 실제 div2=0. 원본 TB가 첫 불일치에서 중단했다. RTL을 원복해 LAB2_PASS clock_divider checks=129를 다시 확인했다.',
        'modified_img':ROOT/'lab2_02_clock_divider/evidence/modified/clock_divider_modified_waveform.png',
        'modified_caption':'수정 파형에서 16 ns 실패 직전 div2는 계속 0이며 tick2가 15 ns에 올라간다.',
        'conclusion':'상태 카운트는 계속 주 클록의 상승 에지에서 갱신한다. 10분주 비교식은 상승 경계를 한 주기 늦춰 60 ns low, 40 ns high가 되며, DIVISOR=2의 검사는 경계 조건이 겹치는 결함을 잡아낸다.',
        'pdf':'05.LAB2_02_CLOCK_DIVIDER_VIVADO.pdf, 13–17쪽, 34–44쪽'
    },
    {
        'id':'03','name':'레지스터 쌍','module':'register_pair','end':'66 ns','checks':'7',
        'goal':'입력값 저장용 stored와 저장값 전달용 value를 분리해, 저장·전달 제어의 우선순위와 동시 동작을 확인한다.',
        'structure':'두 4비트 레지스터는 모두 clk 상승 에지에서 갱신된다. reset은 둘 다 0으로 만든다. load와 transfer는 독립된 조건문이므로 동시에 동작할 수 있다. nonblocking 대입 때문에 같은 에지에서 value가 받는 것은 갱신 전 stored 값이다.',
        'code':"""always @(posedge clk) begin
    if (rst) begin
        stored <= 4'd0;
        value <= 4'd0;
    end else begin
        if (load) stored <= data_in;
        if (transfer) value <= stored;
    end
end""",
        'input':'clk 주기 10 ns. reset 뒤 data_in=A를 load해 stored=A,value=0을 확인한다. data_in을 3으로 바꾼 채 transfer하면 value=A여야 한다. load와 transfer를 같이 켜면 stored는 3으로 갱신되지만 value는 이전 A를 받아야 하며, 다음 전달 에지에는 value=3이 된다. 이후 hold와 reset 우선순위를 확인한다.',
        'normal':'LAB2_PASS register_pair checks=7. 종료 66 ns. reset, load만 수행, 저장값 전달, 동시 동작에서 이전 stored 사용, 다음 에지 전달, hold, reset 우선순위를 확인했다.',
        'normal_img':ROOT/'lab2_03_register/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 VaporView 캡처. 신호 트리에는 stored와 value가 있으나 표시된 파형 행에는 포함되지 않았다.',
        'mod':'교안 지시에 따라 transfer 대입의 우변을 stored에서 data_in으로 바꾸고 TB는 그대로 둔다. 이 수정은 등록된 값을 전달하지 않고 전달 에지의 입력을 직접 value에 저장한다.',
        'failure':'26 ns · transfer stored not live input 검사. 기대 stored=A,value=A, 즉 8’hAA였으나 실제 stored=A,value=3, 즉 8’hA3였다. RTL 원복 후 LAB2_PASS register_pair checks=7을 확인했다.',
        'modified_img':ROOT/'lab2_03_register/evidence/modified/register_pair_modified_waveform.png',
        'modified_caption':'25 ns 전달 에지에서 stored=A와 data_in=3이 함께 보이고, 수정 회로의 value는 3으로 갱신된다.',
        'conclusion':'저장 단계와 전달 단계는 독립된 레지스터 동작이다. 동시 갱신 시 오른쪽 값은 에지 직전의 stored라는 점을 TB가 검사한다. 정상 스크린샷에서 출력 bus가 표시되지 않으므로 수정 파형과 보관 VCD를 함께 해석한다.',
        'pdf':'05.LAB2_03_REGISTER_VIVADO.pdf, 13–14쪽, 35–44쪽'
    },
    {
        'id':'04','name':'4비트 시프트 레지스터','module':'shift_register4','end':'106 ns','checks':'8',
        'goal':'직렬 입력 비트를 최상위 비트에 넣고, 매 enable 에지마다 기존 비트를 하위 방향으로 한 칸 이동한다.',
        'structure':'동기 reset은 값을 0으로 만든다. enable이 1일 때만 {serial_in,value[3:1]}을 저장하므로 serial_in은 bit 3으로 들어가고 기존 bit 3:1은 bit 2:0으로 이동한다. enable=0이면 값은 유지된다.',
        'code':"""always @(posedge clk) begin
    if (rst) value <= 4'd0;
    else if (enable)
        value <= {serial_in, value[3:1]};
end""",
        'input':'TB에서 reset 이후 enable=1로 serial_in=1을 인가하고 15 ns 상승 에지에서 샘플링한다. 입력 1은 value=1000이 되어야 한다. 이어 입력 비트를 이동하며 출력 상태를 검사하고, enable을 내린 구간에서는 값이 변하지 않아야 한다.',
        'normal':'LAB2_PASS shift_register4 checks=8. 종료 106 ns. 입력이 MSB로 들어가고 이후 이동·유지 동작이 TB 기대값을 통과했다.',
        'normal_img':ROOT/'lab2_04_shift_register/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 실행 VaporView 파형과 PASS 로그.',
        'mod':'교안의 수정은 다음 상태 연결을 {serial_in,value[3:1]}에서 {value[2:0],serial_in}으로 바꾸는 것이다. 즉 입력이 LSB로 들어가며 기존 비트는 반대 방향으로 이동한다. TB 기대값은 유지한다.',
        'failure':'16 ns · input enters MSB 검사. 첫 입력 1 뒤 기대값은 1000, 실제 변경 RTL은 0001이다. RTL 원복 후 LAB2_PASS shift_register4 checks=8, 종료 106 ns를 다시 확인했다.',
        'modified_img':ROOT/'lab2_04_shift_register/evidence/modified/waveform.png',
        'modified_caption':'첫 enable 상승 에지에서 serial_in=1이 bit 0에 나타나며 16 ns 검사에서 실패한다.',
        'conclusion':'연결 순서가 직렬 데이터의 유입 방향을 결정한다. 첫 입력만으로 어느 비트에 데이터가 들어갔는지 분명히 구분되며, 수정 실험 로그와 파형이 교안 예상과 일치한다.',
        'pdf':'05.LAB2_04_SHIFT_REGISTER_VIVADO.pdf, RTL·정상 검사 절, 수정 실험 41쪽'
    },
    {
        'id':'05','name':'4비트 PISO 레지스터','module':'piso4','end':'976 ns','checks':'114',
        'goal':'4비트 데이터를 병렬로 적재한 뒤 MSB부터 한 비트씩 직렬 출력한다.',
        'structure':'serial_out은 value[3]에 연결된다. clk 상승 에지에서 reset, load, enable 순서로 동작을 선택한다. load가 우선하며 data_in을 병렬 적재한다. enable은 {value[2:0],1’b0}로 왼쪽 이동하고 빈 LSB를 0으로 채운다.',
        'code':"""assign serial_out = value[3];
always @(posedge clk) begin
    if (rst) value <= 4'd0;
    else if (load) value <= data_in;
    else if (enable)
        value <= {value[2:0], 1'b0};
end""",
        'input':'TB는 데이터를 load한 뒤 네 번 enable해 MSB 우선으로 직렬 비트를 읽는다. 먼저 word=0을 전달한 뒤 word=1을 병렬 적재한다. word=1의 첫 직렬 비트는 bit[3]=0이어야 한다. 총 114개 비교로 reset, load, shift, 출력 순서를 확인한다.',
        'normal':'LAB2_PASS piso4 checks=114. 종료 976 ns. 적재 우선순위와 MSB-first 네 비트 전송을 포함한 검사가 모두 통과했다.',
        'normal_img':ROOT/'lab2_05_piso/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 실행에서 value와 MSB 직렬 출력의 동작을 확인한다.',
        'mod':'교안의 수정은 serial_out을 value[3]에서 value[0]으로 바꾸는 것이다. TB는 MSB 우선 기대값을 유지하므로 LSB 출력은 첫 데이터에서 바로 어긋난다.',
        'failure':'86 ns · MSB first before edge 검사. word=1에서 기대 serial_out=0, 수정 RTL의 실제 serial_out=1이다. 정상 RTL 복구 후 LAB2_PASS piso4 checks=114, 종료 976 ns를 확인했다.',
        'modified_img':ROOT/'lab2_05_piso/evidence/modified/waveform.png',
        'modified_caption':'word=1을 적재한 뒤 value=0001이며 LSB 출력이 1로 나타나 86 ns에 실패한다.',
        'conclusion':'병렬 word의 첫 출력 위치를 지정하는 인덱스는 직렬 전송 순서를 결정한다. LSB와 MSB 차이는 word=0001 하나만으로 검출된다.',
        'pdf':'05.LAB2_05_PISO_VIVADO.pdf, 수정 실험 41쪽, 해설 43쪽'
    },
    {
        'id':'06','name':'Moore 순환 상태기계','module':'moore_cycle','end':'226 ns','checks':'23',
        'goal':'enable과 advance가 함께 활성화될 때 S0→S1→S2→S0 순환하는 2비트 Moore 상태기계를 구성한다.',
        'structure':'상태 value는 동기 reset에서 00으로 초기화된다. 두 제어가 모두 1인 상승 에지에서만 다음 상태로 간다. 나머지 입력에서는 현재 상태를 보존한다. 출력은 등록된 상태 자체이므로 입력 변화만으로 즉시 바뀌지 않는다.',
        'code':"""always @(posedge clk) begin
    if (rst) value <= 2'b00;
    else if (enable && advance) begin
        case (value)
            2'b00: value <= 2'b01;
            2'b01: value <= 2'b10;
            default: value <= 2'b00;
        endcase
    end
end""",
        'input':'TB는 reset으로 S0=00을 만든다. enable과 advance를 이용해 S0에서 S1, S1에서 S2, S2에서 S0로 진행하는 경우와 정지 조건을 검사한다. 출력 상태는 유효한 상승 에지 뒤에 바뀐다.',
        'normal':'LAB2_PASS moore_cycle checks=23. 종료 226 ns. 상태 전이와 정지 조건에 대한 23개 검사가 통과했다.',
        'normal_img':ROOT/'lab2_06_moore/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 상태 파형에서 상승 에지마다 상태가 순환한다.',
        'mod':'교안 지시에 따라 S1=01의 다음 상태를 S2=10에서 S0=00으로 변경한다. TB의 S1→S2 기대값은 그대로 둔다.',
        'failure':'46 ns · S1 to S2 검사. enable=1, advance=1 에지 뒤 기대값 10, 실제값 00이다. 원복 뒤 LAB2_PASS moore_cycle checks=23, 종료 226 ns를 확인했다.',
        'modified_img':ROOT/'lab2_06_moore/evidence/modified/waveform.png',
        'modified_caption':'S1에서 수정 전이값 00이 출력되어 46 ns 검사에서 S2 기대값과 다르다.',
        'conclusion':'Moore 기계의 다음 상태는 현재 상태와 제어 입력에 따라 결정되며 출력 변화는 클록 경계에서 발생한다. 한 전이만 잘못된 경우에도 해당 경로를 지정한 검사에서 검출된다.',
        'pdf':'05.LAB2_06_MOORE_VIVADO.pdf, 수정 실험 39쪽, 파형 해설 41쪽'
    },
    {
        'id':'07','name':'Mealy 토글 상태기계','module':'mealy_toggle','end':'67 ns','checks':'11',
        'goal':'입력이 1일 때 enable에 따라 상태를 토글하고, 현재 상태와 입력을 함께 이용해 조합 출력을 생성한다.',
        'structure':'state는 상승 에지에서만 갱신된다. reset은 S0을 선택하며 enable과 bit_in이 모두 1일 때 상태가 반전된다. value는 조합 논리다. bit_in=0이면 00, bit_in=1이면 현재 state에 따라 10 또는 01이 된다.',
        'code':"""always @(posedge clk) begin
    if (rst) state <= 1'b0;
    else if (enable && bit_in) state <= ~state;
end
assign value = bit_in ?
    (state ? 2'b01 : 2'b10) : 2'b00;""",
        'input':'TB는 S0을 reset한 뒤 state가 그대로인 동안 bit_in을 0에서 1로 바꾼다. 이 경우 조합 출력은 클록 사이에도 변해야 한다. 이후 enable과 bit_in이 함께 1인 상승 에지에서 상태가 토글되는지 확인한다.',
        'normal':'LAB2_PASS mealy_toggle checks=11. 종료 67 ns. 상태 토글, 입력 의존 조합 출력, 클록 사이 출력 변화를 포함한 11개 검사가 통과했다.',
        'normal_img':ROOT/'lab2_07_mealy/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 파형에서 입력 변화가 value를 즉시 바꾸고 상태는 상승 에지에서 갱신된다.',
        'mod':'교안은 bit_in=1일 때 S0과 S1의 출력을 서로 바꾸도록 한다. 원래 S0 출력은 10, S1 출력은 01이며 TB 기대값을 유지한다.',
        'failure':'7 ns · S0 input changes between clocks 검사. state=0, bit_in=1, enable=0에서 기대 {state,value}=010, 실제 001이다. RTL 복구 후 LAB2_PASS mealy_toggle checks=11, 종료 67 ns를 확인했다.',
        'modified_img':ROOT/'lab2_07_mealy/evidence/modified/waveform.png',
        'modified_caption':'S0에서 입력이 1로 바뀐 직후 수정 회로의 조합 출력 01이 나타나 7 ns에 실패한다.',
        'conclusion':'상태는 순차 논리, 출력은 조합 논리로 구현되어 있어 입력 변화와 상태 변화의 시점이 다르다. 클록 사이 검사는 Mealy 출력의 입력 의존성을 확인한다.',
        'pdf':'05.LAB2_07_MEALY_VIVADO.pdf, 수정 실험 41쪽, 파형 해설 43쪽'
    },
    {
        'id':'08','name':'8자리 7세그먼트 스캔','module':'segment_scan8','end':'1306 ns','checks':'194',
        'goal':'8자리 표시기를 빠르게 순차 선택하고 각 4비트 숫자에 맞는 세그먼트 패턴을 출력한다.',
        'structure':'index가 현재 자리를 가리키고, digits에서 해당 4비트를 골라 16진수 디코더로 segments를 만든다. 한 자리 전환 동안 blank를 활성화해 모든 select를 0으로 끄고, 다음 구간에서 one-hot 자리를 켠다. 보드 polarity는 wrapper에서 적용한다.',
        'code':"""if (rst) begin
    index <= 0;
    blank <= 1'b1;
end else if (enable) begin
    if (blank) blank <= 0;
    else begin
        blank <= 1;
        index <= index + 1'b1;
    end
end
nibble = digits >> (index * 4);
select = blank ? 8'h00 : (8'b0000_0001 << index);""",
        'input':'TB는 reset 동안 index=0, select=00으로 모든 자리가 꺼지는지 확인한다. 이후 enable로 blank와 one-hot select가 번갈아 동작하고, 자리 index와 digits의 4비트 nibble에 맞는 segments 패턴이 나오는지 검사한다.',
        'normal':'LAB2_PASS segment_scan8 checks=194. 종료 1306 ns. 리셋 blanking, 자리 선택, index 순환, 숫자 디코딩 검사가 통과했다.',
        'normal_img':ROOT/'lab2_08_segment_scan/evidence/normal/normal_screenshot.png',
        'normal_caption':'정상 VaporView 파형에서 blank 구간과 one-hot 자리 선택, 디코더 출력이 교대한다.',
        'mod':'교안에 따라 선택 신호의 blank 조건을 제거하고 one-hot 선택만 유지한다. TB는 변경하지 않아 reset 동안 모든 자리가 꺼져야 한다는 기대를 유지한다.',
        'failure':'6 ns · reset blanks digit zero 검사. 기대 select=00, 실제 select=01이며 index=0에서 첫 자리가 켜졌다. RTL 원복 후 LAB2_PASS segment_scan8 checks=194, 종료 1306 ns를 확인했다.',
        'modified_img':ROOT/'lab2_08_segment_scan/evidence/modified/waveform.png',
        'modified_caption':'리셋 중 index=0, select=01 및 숫자 0 패턴 FC가 보이며 6 ns에 실패한다.',
        'conclusion':'자리 전환 blanking은 표시값이 바뀌는 구간에 이전 자리 또는 엉뚱한 세그먼트가 잠깐 켜지는 현상을 막는다. reset blanking 검사는 표시기 안전 상태를 첫 시점에 확인한다.',
        'pdf':'05.LAB2_08_SEGMENT_SCAN_VIVADO.pdf, 수정 실험 48쪽, 파형 해설 50쪽'
    },
]

MOD_EXPECTED_ACTUAL = {
    '02': ('div2=1', 'div2=0'),
    '03': ("stored=A,value=A (8'hAA)", "stored=A,value=3 (8'hA3)"),
    '04': ('value=1000', 'value=0001'),
    '05': ('serial_out=0', 'serial_out=1'),
    '06': ('value=10', 'value=00'),
    '07': ('{state,value}=010', '{state,value}=001'),
    '08': ('select=00', 'select=01'),
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '5')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), GRAY)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_section_columns(section, count=2, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sect_pr.append(cols)
    cols.set(qn('w:num'), str(count))
    cols.set(qn('w:space'), str(space_twips))
    cols.set(qn('w:equalWidth'), '1')


def crop_image(src, dst, box):
    dst.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as im:
        im.crop(box).save(dst)
    return dst


def prepare_evidence_images():
    """Create readable crops for the narrow two-column body."""
    assets = {}
    split_ratios = {
        '01': .58, '02': .59, '03': .59, '04': .59,
        '05': .61, '06': .62, '07': .61, '08': .62,
    }
    for lab in LABS:
        lab_id = lab['id']
        vsc = (ROOT / 'lab2_01_counter/evidence/normal/normal_screenshot.png'
               if lab_id == '01' else lab['normal_img'])
        with Image.open(vsc) as im:
            w, h = im.size
        wave = ASSET_DIR / f'lab2_{lab_id}_normal_waveform_ui.png'
        terminal = ASSET_DIR / f'lab2_{lab_id}_normal_terminal.png'
        crop_image(vsc, wave, (0, 0, w, int(h * .42)))
        crop_image(vsc, terminal, (0, int(h * split_ratios[lab_id]), w, h))
        assets[(lab_id, 'normal_wave')] = wave
        assets[(lab_id, 'normal_terminal')] = terminal

        modified = lab['modified_img']
        if modified:
            with Image.open(modified) as im:
                mw, mh = im.size
            if mw >= 3000:
                left = ASSET_DIR / f'lab2_{lab_id}_modified_a.png'
                right = ASSET_DIR / f'lab2_{lab_id}_modified_b.png'
                crop_image(modified, left, (0, 0, int(mw * .56), mh))
                crop_image(modified, right, (int(mw * .44), 0, mw, mh))
                assets[(lab_id, 'modified')] = [left, right]
            else:
                assets[(lab_id, 'modified')] = [modified]

    counter = LABS[0]['normal_img']
    with Image.open(counter) as im:
        cw, ch = im.size
    counter_overview = ASSET_DIR / 'lab2_01_counter_overview.png'
    counter_details = ASSET_DIR / 'lab2_01_counter_details.png'
    crop_image(counter, counter_overview, (0, 0, cw, int(ch * .55)))
    crop_image(counter, counter_details, (0, int(ch * .47), cw, ch))
    assets[('01', 'overview')] = counter_overview
    assets[('01', 'details')] = counter_details
    return assets


def set_run_font(run, font=FONT_BODY, size=10.5, bold=False, color='000000'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn('w:' + key), font)


def style_para(p, before=0, after=6, line=1.16, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_together = True
    if keep:
        pf.keep_with_next = True


def add_text(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    style_para(p)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    r = p.add_run(text)
    set_run_font(r, FONT_HEAD, 15 if level == 1 else 11.5, bold=True)
    return p


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    table.columns[0].width = Inches(COLUMN_WIDTH)
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'F3F5F7')
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    style_para(p, after=0, line=1.05)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, FONT_CODE, 7.4)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    if widths:
        scale = min(1.0, COLUMN_WIDTH / sum(widths))
        widths = [width * scale for width in widths]
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    set_repeat_table_header(table.rows[0])
    for i, txt in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=115, bottom=115, start=125, end=125)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_para(p, after=0, line=1.05)
        r = p.add_run(str(txt))
        set_run_font(r, FONT_HEAD, font_size, bold=True, color='FFFFFF')
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            if widths:
                cell.width = Inches(widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=105, bottom=105, start=125, end=125)
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            style_para(p, after=0, line=1.08)
            if cidx == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, FONT_BODY, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def set_alt_text(shape, title, descr):
    inline = shape._inline
    docpr = inline.docPr
    docpr.set('title', title)
    docpr.set('descr', descr)


def add_figure(doc, path, caption, number, width=3.02):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, f'그림 {number}', caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True
    style_para(cap, after=7, line=1.0)
    r = cap.add_run(f'그림 {number}. {caption}')
    set_run_font(r, FONT_BODY, 9.1)


def add_figure_pair(doc, left_path, left_caption, left_number, right_path, right_caption, right_number):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    set_table_borders(table)
    for col in table.columns:
        col.width = Inches(3.18)
    for cell, path in zip(table.rows[0].cells, [left_path, right_path]):
        cell.width = Inches(3.18)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        shape = p.add_run().add_picture(str(path), width=Inches(3.0))
        set_alt_text(shape, 'LAB2-01 정상 증거', '정상 시뮬레이션 결과')
    for cell, caption, number in zip(table.rows[1].cells, [left_caption, right_caption], [left_number, right_number]):
        cell.width = Inches(3.18)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=30, bottom=60, start=75, end=75)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_para(p, after=0, line=1.0)
        r = p.add_run(f'그림 {number}. {caption}')
        set_run_font(r, FONT_BODY, 8.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, FONT_BODY, 9, color='555555')
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t'); txt.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, sep, txt, end])


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    assets = prepare_evidence_images()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.72)
    sec.left_margin = Inches(.82)
    sec.right_margin = Inches(.82)
    sec.footer_distance = Inches(.38)

    normal = doc.styles['Normal']
    normal.font.name = FONT_BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)
    for name, size in [('Title',28),('Subtitle',14),('Heading 1',15),('Heading 2',11.5)]:
        st = doc.styles[name]
        st.font.name = FONT_HEAD if name != 'Subtitle' else FONT_BODY
        st.font.size = Pt(size)
        st.font.bold = name != 'Subtitle'
        st.font.color.rgb = RGBColor(0,0,0)
        st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_HEAD if name != 'Subtitle' else FONT_BODY)
        if name == 'Title' and st._element.pPr is not None:
            pborder = st._element.pPr.find(qn('w:pBdr'))
            if pborder is not None:
                st._element.pPr.remove(pborder)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_para(footer, after=0, line=1.0)
    r = footer.add_run('LAB2 실험 전 보고서  ·  ')
    set_run_font(r, FONT_BODY, 9, color='555555')
    add_page_field(footer)

    # Cover
    for _ in range(5):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(13)
    r = p.add_run('LAB2 실험 전 보고서'); set_run_font(r, FONT_HEAD, 28, bold=True)
    p = doc.add_paragraph(style='Subtitle')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run('순차 회로와 상태 기계 사전 시뮬레이션'); set_run_font(r, FONT_BODY, 14)
    for key, value in [('과목','전자회로설계실험 2'),('제출일','2026년 9월 21일'),('성명','____________________________'),('학번','____________________________')]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(13)
        r = p.add_run(f'{key}  {value}'); set_run_font(r, FONT_BODY, 12)
    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec.page_width = Inches(8.5)
    body_sec.page_height = Inches(11)
    body_sec.top_margin = Inches(.72)
    body_sec.bottom_margin = Inches(.72)
    body_sec.left_margin = Inches(.82)
    body_sec.right_margin = Inches(.82)
    body_sec.footer_distance = Inches(.38)
    set_section_columns(body_sec, 2, 360)

    add_heading(doc, '1 보고서 범위와 결론', 1)
    add_text(doc, '이 보고서는 LAB2의 필수 실험 01–08을 대상으로 정상 회로의 기능, 입력 조건, 자기검사 결과와 교안에서 지정한 수정 실험을 정리한다. 여덟 정상 설계 모두 기록된 PASS와 종료 시각이 확인되었다. LAB2-02부터 LAB2-08까지는 교안이 지정한 한 곳을 바꿔 원본 테스트벤치에서 예상한 첫 FAIL을 확인했고, RTL 복구 뒤 정상 PASS를 다시 확인했다. LAB2-01 교안에는 별도 수정 실험이 없어 고장 주입을 하지 않았다.')
    add_text(doc, 'PASS는 기록된 테스트벤치가 지정한 비교를 통과했다는 뜻이다. 보드 핀, 입력 전처리기, FPGA 구현이나 모든 가능한 입력 조합까지 검증했다는 뜻은 아니다. 실험 파일은 2026년 9월 21일 로컬 작업 트리 스냅샷이므로, 제출 전 변경 파일을 commit하고 제출용 tag를 붙여 재현 근거를 고정해야 한다.')
    add_table(doc, ['실험','검증 요약'], [[f"{x['id']} {x['name']}",
        f"정상 {x['checks']}회 · 종료 {x['end']} · " + ('수정 없음' if x['id']=='01' else x['failure'])]
        for x in LABS], [.88,2.24], 8.1)

    doc.add_page_break()
    add_heading(doc, '2 실험 환경과 검증 방법', 1)
    add_text(doc, '코어 RTL은 Verilog로 작성하고 VS Code의 실행 작업에서 Icarus Verilog(iverilog/vvp)로 컴파일·시뮬레이션했다. 테스트벤치는 VCD를 저장하고 VaporView에서 clk, reset, 제어 입력, 상태 또는 출력 파형을 비교했다. 보드 top과 XDC는 프로젝트에 포함되어 있지만 이 코어 TB는 XDC를 읽지 않는다. 따라서 PASS는 핀 할당과 실제 보드 입출력의 검증을 포함하지 않는다.')
    add_text(doc, 'TB의 clk 주기는 10 ns로 설정되어 있다. 순차 로직은 posedge clk에서 nonblocking assignment로 다음 상태를 예약하고, 검사 task는 에지 뒤 1 ns 기다린 후 출력을 비교한다. 이 시간 간격은 같은 에지에서 여러 레지스터를 갱신할 때 기존 상태와 새 상태를 구분해 관찰하도록 한다. 실험별 예외나 입력 순서는 각 장에 적었다.')
    add_heading(doc, '3 소스 출처와 실행 스냅샷', 1)
    add_table(doc, ['항목','식별 정보'], [
        ['템플릿 원격','https://github.com/Glaysia/fpga-lab-template.git'],
        ['기준 태그·커밋','v2.0.1 · 86c15c556e2edf1564a203180690f7ec4ad056af'],
        ['상위 저장소 remote','https://github.com/tae2yi/UOS_ece2.git'],
        ['상위 저장소 HEAD','2609ab4b461c5add334791ebf8d69a721434d9fc'],
    ], [1.65,5.1], 9.0)
    add_text(doc, '템플릿 주소와 태그는 코드의 기준 출처다. 실험별 source, 테스트벤치와 evidence가 기준 태그 상태와 같다고 간주하지 않으며, 현재 보고된 결과는 각 실험 폴더의 normal/modified/recovered 로그와 VCD에 연결한다.')

    fig_no = 1
    for x in LABS:
        add_heading(doc, f"LAB2-{x['id']} {x['name']}", 1)
        add_heading(doc, '설계 목표와 회로 구조', 2)
        add_text(doc, x['goal'])
        add_text(doc, x['structure'])
        add_heading(doc, '핵심 코드', 2)
        add_text(doc, f"정상 코어 {x['module']}의 상태 갱신 또는 출력 생성 부분이다. 보드 연결 모듈과 입력 전처리 코드는 생략했다.")
        add_code(doc, x['code'])
        add_heading(doc, '입력 조건과 예상 결과', 2)
        add_text(doc, x['input'])
        add_heading(doc, '정상 시뮬레이션 결과', 2)
        add_text(doc, x['normal'])
        if x['id'] == '01':
            add_figure(doc, assets[('01', 'overview')], '정상 전체 구간의 증가·감소와 wrap.', fig_no); fig_no += 1
            add_figure(doc, assets[('01', 'details')], 'wrap, hold 및 동기 reset 상세 구간.', fig_no); fig_no += 1
        else:
            add_figure(doc, assets[(x['id'], 'normal_wave')], x['normal_caption'], fig_no); fig_no += 1
        add_figure(doc, assets[(x['id'], 'normal_terminal')],
                   f"LAB2-{x['id']} 정상 실행의 PASS 및 종료 시각 터미널 로그.", fig_no); fig_no += 1
        add_heading(doc, '수정 실험 및 예상 오류', 2)
        add_text(doc, x['mod'])
        add_heading(doc, '수정 실행 결과', 2)
        if x['id']=='01':
            add_text(doc, x['failure'])
        else:
            expected, actual = MOD_EXPECTED_ACTUAL[x['id']]
            check_at = x['failure'].split(' · ')[0]+' · '+x['failure'].split(' · ')[1].split('.')[0]
            add_table(doc, ['구분','내용'], [['시점·검사', check_at], ['기대', expected], ['실제', actual]], [.78,2.34], 8.4)
            add_text(doc, x['failure'])
            modified_parts = assets[(x['id'], 'modified')]
            if len(modified_parts) == 1:
                add_figure(doc, modified_parts[0], x['modified_caption'], fig_no); fig_no += 1
            else:
                add_figure(doc, modified_parts[0], x['modified_caption'] + ' 전반 구간.', fig_no); fig_no += 1
                add_figure(doc, modified_parts[1], x['modified_caption'] + ' 실패 시점 확대 구간.', fig_no); fig_no += 1
        add_heading(doc, '문제 해결 과정과 결론', 2)
        add_text(doc, x['conclusion'])

    add_heading(doc, '9 비교와 결론', 1)
    add_text(doc, '여덟 실험은 상승 에지에서 상태가 바뀌는 순차 회로와 입력 변화에 따라 조합 출력이 바뀌는 Mealy 회로를 함께 다룬다. reset 우선순위, enable에 의한 유지, nonblocking assignment의 이전 상태 참조, 카운터 경계, 데이터 경로 방향, 상태 전이, 자리 blanking이 각각 독립된 검사로 확인되었다.')
    add_table(doc, ['실험','핵심 관찰과 결과'], [
        ['01 카운터','방향·enable·wrap · 정상 36 / 356 ns · 수정 없음'],
        ['02 분주기','주기·duty·tick · 정상 129 / 436 ns · 16 ns, div2 1→0'],
        ['03 레지스터','저장값·이전 상태 · 정상 7 / 66 ns · 26 ns, AA→A3'],
        ['04 시프트','직렬 입력 방향 · 정상 8 / 106 ns · 16 ns, 1000→0001'],
        ['05 PISO','MSB-first · 정상 114 / 976 ns · 86 ns, 0→1'],
        ['06 Moore','상태 순환 · 정상 23 / 226 ns · 46 ns, 10→00'],
        ['07 Mealy','입력 의존 출력 · 정상 11 / 67 ns · 7 ns, 010→001'],
        ['08 자리 스캔','blank·one-hot · 정상 194 / 1306 ns · 6 ns, 00→01'],
    ], [.84,2.28], 8.1)
    add_text(doc, '각 수정 실험에서는 한 동작만 바꾸고 테스트벤치를 유지했다. 따라서 첫 FAIL 시점은 변경된 동작과 기대값이 처음 갈라지는 지점이다. 정상 VCD와 수정 VCD를 함께 비교하고 RTL을 원복한 재실행까지 확인함으로써 의도된 오류와 복구 상태를 구분했다. 08A 통합 실습은 공통 프로젝트 흐름에만 언급되며 필수 실험 수에는 포함하지 않았다.')

    add_heading(doc, '참고 문헌', 1)
    add_text(doc, 'M. Morris Mano and Michael D. Ciletti, Digital Design: With an Introduction to the Verilog HDL, VHDL, and SystemVerilog, 6th ed., Pearson, 2018, ISBN 978-0-13-454989-7.')
    doc.core_properties.title = 'LAB2 실험 전 보고서'
    doc.core_properties.subject = '순차 회로와 상태 기계 사전 시뮬레이션'
    doc.core_properties.author = ''
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
