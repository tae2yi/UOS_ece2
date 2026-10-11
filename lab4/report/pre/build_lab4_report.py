"""LAB4 실험 전(예비) 보고서 생성기 — LAB3 예비보고서의 레이아웃(A4, 2단)을 따른다.

대상: STM32_0(LED·버튼 폴링, 과제1·2)과 FPGA 교재 실습 2개(4비트 곱셈기, BCD→이진수 변환).
그림은 make_figures.py가 먼저 만든다.

실행: python build_lab4_report.py [커밋 해시]   (python-docx, Pillow 필요)
      커밋 해시를 주면 보고서의 GitHub 주소가 그 커밋을 가리킨다. 생략하면 master.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
IMG = HERE / 'images'
OUT = HERE / 'LAB4_예비보고서.docx'

FONT = 'Nanum Gothic'
FONT_CODE = 'Consolas'
NAVY = '1F4E78'
PALE = 'EAF2F8'
GRAY = 'D9D9D9'
REPO = 'https://github.com/tae2yi/UOS_ece2'
CODE_COMMIT = sys.argv[1] if len(sys.argv) > 1 else 'master'
UPSTREAM = 'https://github.com/Glaysia/ece2_474'
UPSTREAM_COMMIT = 'dd94f01047aebdece09b9620c66e9b179403e5c9'
COL_W = 3.22          # 2단 본문 한 단의 사용 폭(inch)
FIG_W = 3.18
CODE_MAX = 63         # 한 단에 들어가는 코드 한 줄의 최대 글자 수(6.4 pt)

# ============================================================ docx helpers
def gh(path, first=None, last=None, tree=False):
    url = f"{REPO}/{'tree' if tree else 'blob'}/{CODE_COMMIT}/{path}"
    return url + (f'#L{first}-L{last}' if first else '')


def add_hyperlink(paragraph, url, text, size=7.0, color='0563C1'):
    r_id = paragraph.part.relate_to(
        url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), FONT)
    rpr.append(fonts)
    c = OxmlElement('w:color'); c.set(qn('w:val'), color); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rpr.append(sz)
    run.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_link_line(doc, label, url):
    p = doc.add_paragraph()
    para_fmt(p, before=1, after=5, line=1.0)
    set_font(p.add_run(f'{label}: '), size=7.0, bold=True, color='555555')
    add_hyperlink(p, url, url)
    return p


def set_font(run, name=FONT, size=10.0, bold=False, color='000000', italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), name)


def set_style_font(style, name, size, bold=None, color='000000'):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), name)


def para_fmt(p, before=0, after=4, line=1.12, keep_next=False, keep_together=True, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_together = keep_together
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None:
        p.alignment = align


def add_text(doc, text, size=10.0, after=4, keep_next=False):
    text = text.replace('‥', '~')
    p = doc.add_paragraph()
    para_fmt(p, after=after, keep_next=keep_next)
    set_font(p.add_run(text), size=size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    para_fmt(p, before=10 if level == 1 else 7, after=5 if level == 1 else 3, line=1.0, keep_next=True)
    set_font(p.add_run(text), size=15.0 if level == 1 else 11.5, bold=True)
    return p


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(round(inches * 1440)))
    tc_w.set(qn('w:type'), 'dxa')


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)


def table_setup(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(round(sum(widths) * 1440)))
    tbl_w.set(qn('w:type'), 'dxa')
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), GRAY)
        borders.append(el)
    tbl_pr.append(borders)
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', 45), ('left', 70), ('bottom', 45), ('right', 70)):
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tbl_pr.append(mar)
    for col, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        col.set(qn('w:w'), str(round(w * 1440)))


def row_props(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))
    if header:
        h = OxmlElement('w:tblHeader')
        h.set(qn('w:val'), 'true')
        tr_pr.append(h)


def spacer(doc, pt=3):
    p = doc.add_paragraph()
    para_fmt(p, after=0, line=1.0)
    p.paragraph_format.space_after = Pt(pt)
    r = p.add_run()
    set_font(r, size=2)
    return p


def add_table(doc, headers, rows, widths, size=7.4, center_first=True, links=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table_setup(table, widths)
    row_props(table.rows[0], header=True)
    for i, txt in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        para_fmt(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(txt), size=size + 0.4, bold=True, color='FFFFFF')
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        row_props(table.rows[-1])
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            set_cell_width(cell, widths[cidx])
            set_cell_shading(cell, PALE if ridx % 2 == 0 else 'FFFFFF')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if (cidx == 0 and center_first) else WD_ALIGN_PARAGRAPH.LEFT
            para_fmt(p, after=0, line=1.0, align=align)
            if links and cidx == 0 and links[ridx]:
                add_hyperlink(p, links[ridx], str(value), size=size)
            else:
                set_font(p.add_run(str(value).replace('‥', '~')), size=size, color='222222')
    spacer(doc, 4)
    return table


def add_code(doc, code, size=6.8, link=None, label='GitHub'):
    lines = []
    for line in code.splitlines():
        body = line.lstrip(' ')
        lines.append(' ' * ((len(line) - len(body)) // 2) + body)
    too_long = [l for l in lines if len(l) > CODE_MAX]
    if too_long:
        raise ValueError(f'code line too long for one column: {too_long}')
    code = '\n'.join(lines)
    table = doc.add_table(rows=1, cols=1)
    table_setup(table, [COL_W])
    row_props(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_width(cell, COL_W)
    set_cell_shading(cell, 'F2F2F2')
    p = cell.paragraphs[0]
    para_fmt(p, after=0, line=1.0)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(line), name=FONT_CODE, size=size, color='222222')
    if link:
        add_link_line(doc, label, link)
    else:
        spacer(doc, 4)


def add_figure(doc, path, caption, number, width=FIG_W):
    p = doc.add_paragraph()
    para_fmt(p, before=4, after=1, line=1.0, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    docpr = shape._inline.docPr
    docpr.set('title', f'그림 {number}')
    docpr.set('descr', caption)
    cap = doc.add_paragraph(style='Caption')
    para_fmt(cap, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(cap.add_run(f'그림 {number}. {caption}'), size=8.0, bold=True, color='333333')
    return number + 1


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_font(run, size=7.0, color='7F7F7F')
    for tag, attr in (('w:fldChar', 'begin'), ('w:instrText', None), ('w:fldChar', 'separate'),
                      ('w:t', None), ('w:fldChar', 'end')):
        el = OxmlElement(tag)
        if tag == 'w:fldChar':
            el.set(qn('w:fldCharType'), attr)
        elif tag == 'w:instrText':
            el.set(qn('xml:space'), 'preserve')
            el.text = ' PAGE '
        else:
            el.text = '1'
        run._r.append(el)


def column_break(doc):
    p = doc.add_paragraph()
    para_fmt(p, after=0, line=1.0)
    p.add_run().add_break(WD_BREAK.COLUMN)


def page_setup(section, columns):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.header_distance = section.footer_distance = Cm(0.8)
    cols = section._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        section._sectPr.append(cols)
    cols.set(qn('w:num'), str(columns))
    cols.set(qn('w:space'), '425')


def bullets(doc, items, size=9.5):
    for item in items:
        p = doc.add_paragraph()
        para_fmt(p, after=2)
        p.paragraph_format.left_indent = Inches(0.14)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        set_font(p.add_run('· ' + item), size=size)
    spacer(doc, 2)


def log_lines(name, keep):
    text = (HERE / 'logs' / name).read_text(encoding='utf-8', errors='replace')
    return [line.rstrip() for line in text.splitlines() if any(k in line for k in keep)]


# ============================================================ 본문에 싣는 코드
TASK1_CODE = """void polling_routine(void)
{
    static GPIO_PinState previous = GPIO_PIN_RESET;
    static uint32_t last_press_ms = 0;
    static uint32_t last_toggle_ms = 0;
    static uint32_t interval_ms = 1000U;
    GPIO_PinState current =
        HAL_GPIO_ReadPin(BUTTON_ece2_GPIO_Port, BUTTON_ece2_Pin);
    uint32_t now_ms = HAL_GetTick();

    if (current == GPIO_PIN_SET && previous == GPIO_PIN_RESET
        && now_ms - last_press_ms >= 200U)
    {
        last_press_ms = now_ms;
        interval_ms = (interval_ms >= 8000U)
                      ? 1000U : interval_ms * 2U;
    }
    previous = current;

    if (now_ms - last_toggle_ms >= interval_ms)
    {
        last_toggle_ms = now_ms;
        HAL_GPIO_TogglePin(BLINK_LED_GPIO_Port, BLINK_LED_Pin);
    }
}"""

TASK2_CODE = """#ifndef STM32_0_USE_EXTERNAL_LED
#define STM32_0_USE_EXTERNAL_LED 0
#endif

#if STM32_0_USE_EXTERNAL_LED
#define BLINK_LED_GPIO_Port GPIO_OUT1_ece2_GPIO_Port
#define BLINK_LED_Pin GPIO_OUT1_ece2_Pin
#else
#define BLINK_LED_GPIO_Port LED_ece2_GPIO_Port
#define BLINK_LED_Pin LED_ece2_Pin
#endif"""

GUARD_CODE = """if (current != previous)
{
    if (current == GPIO_PIN_SET
        && now_ms - last_change_ms >= 50U)
    {
        interval_ms = (interval_ms >= 8000U)
                      ? 1000U : interval_ms * 2U;
    }
    last_change_ms = now_ms;
    previous = current;
}"""

MULT_CODE = """module mult_4bit(input wire [3:0] a, b, output wire [7:0] m);
    multiply_unsigned #(.WIDTH(4)) u_multiply(a, b, m);
endmodule

module multiply_unsigned #(parameter WIDTH = 4)(
    input wire [WIDTH-1:0] a, b,
    output wire [2*WIDTH-1:0] product
);
    assign product = a * b;
endmodule"""

BCD_CODE = """module bcd_to_binary(
    input wire [3:0] ten, one,
    output wire [6:0] binary,
    output wire valid
);
    assign binary = ten * 7'd10 + one;
    assign valid = (ten <= 9) && (one <= 9);
endmodule"""


# ============================================================ build
def build():
    doc = Document()
    page_setup(doc.sections[0], 1)

    set_style_font(doc.styles['Normal'], FONT, 10.0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    for name, size, bold in (('Title', 25.0, True), ('Subtitle', 12.0, False),
                             ('Heading 1', 15.0, True), ('Heading 2', 11.5, True), ('Caption', 8.0, True)):
        set_style_font(doc.styles[name], FONT, size, bold, '333333' if name == 'Caption' else '000000')
    title_ppr = doc.styles['Title'].element.pPr
    if title_ppr is not None and title_ppr.find(qn('w:pBdr')) is not None:
        title_ppr.remove(title_ppr.find(qn('w:pBdr')))

    # ---- cover
    p = doc.add_paragraph(style='Title')
    para_fmt(p, before=150, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run('LAB4 예비 보고서'), size=25.0)
    p = doc.add_paragraph(style='Subtitle')
    para_fmt(p, after=34, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run('STM32_0 LED·버튼 폴링 — FPGA 4비트 곱셈기·BCD→이진수 변환'), size=12.0, italic=True)
    for key, value in (('과목', '전자전기컴퓨터설계실험 2'), ('제출일', '2026년 10월 12일'),
                       ('성명', '김태이'), ('학번', '2023440158')):
        p = doc.add_paragraph()
        para_fmt(p, after=7, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(f'{key}  {value}'), size=10.0)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(body, 2)
    for sec in doc.sections:
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        para_fmt(fp, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(fp.add_run('LAB4 실험 전 보고서  ·  '), size=7.0, color='7F7F7F')
        add_page_field(fp)

    fig = 1
    lab4 = 'lab4'
    routine = f'{lab4}/stm32_lab/Core/Src/ece2_474_routine.c'

    # ---- 1 범위와 결론
    add_heading(doc, '1 보고서 범위와 결론', 1)
    add_text(doc, '10월 12일 실험(STM32_0)은 두 부분이다. NUCLEO-G474RE에서는 제공 예제 2개(1초 간격 LED, 버튼으로 LED 변경)를 실행한 뒤 '
                  '과제1(버튼으로 점멸 간격을 1→2→4→8→1초로 변경)과 과제2(같은 동작을 외부 LED로 구현)를 직접 작성한다. '
                  'Combo II-DLD에서는 교재 실습 4비트 곱셈기와 BCD→이진수 변환을 시뮬레이션하고 보드에서 시연한다[1].')
    add_text(doc, '이 보고서는 실험 전에 설계 목표와 방법, 핵심 코드와 회로 구조, 사전 검증 결과, 보드에서 확인할 항목을 정리한다. '
                  '과제1·2의 코드는 미리 작성해 Arm GCC로 빌드했고, 같은 소스 파일을 PC에서 HAL 대역과 함께 실행해 '
                  '간격 순환, 길게 누름, 접점 떨림, tick 넘침의 네 경우를 확인했다. '
                  'FPGA 두 설계는 Icarus Verilog와 Vivado 시뮬레이터에서 256개 입력 조합을 모두 통과했고, '
                  'Vivado에서 bitstream 생성까지 마쳤다.')
    add_text(doc, '사전 검증에서 한 가지 위험을 찾았다. 제공 예제의 200 ms 잠금은 누르는 순간의 떨림만 막으므로, '
                  '200 ms보다 길게 눌렀다 뗄 때 떨림이 있으면 간격이 한 번 더 바뀐다. 실제 버튼에서 이 현상이 나오는지는 '
                  '보드에서 확인할 항목이며, 나올 때 바꿔 넣을 대비안을 준비했다.')
    column_break(doc)
    add_table(doc, ['항목', '사전 검증', '결과'], [
        ['제공 예제 1·2', '호스트 시뮬레이션', '1초 간격, 누름당 1회'],
        ['과제1', 'Arm GCC 빌드, 호스트 4개 시나리오', '1→2→4→8→1초'],
        ['과제2', 'Arm GCC 컴파일, 호스트', 'PB0에서 과제1과 같은 시각'],
        ['곱셈기', 'Icarus·Vivado 256쌍', 'PASS, LUT 16개'],
        ['BCD 변환', 'Icarus·Vivado 256개', 'PASS, LUT 7개'],
    ], [0.82, 1.28, 1.12], size=7.2)
    add_text(doc, '여기서 PASS와 호스트 시뮬레이션 결과는 논리가 맞다는 뜻이다. 핀 결선, LED·스위치의 극성, 실제 버튼의 떨림은 '
                  '포함하지 않으므로 각 장 끝의 표대로 보드에서 따로 확인한다.', size=8.5)

    # ---- 2 환경
    add_heading(doc, '2 실험 환경과 검증 방법', 1)
    add_table(doc, ['구분', '사용한 것'], [
        ['MCU 보드', 'NUCLEO-G474RE (STM32G474RE, Cortex-M4)'],
        ['MCU 프로젝트', 'Glaysia/ece2_474 v1.1.0'],
        ['MCU 빌드', 'Arm GNU Toolchain 13.3.Rel1, CMake 3.24.2, Ninja 1.11.1'],
        ['호스트 검증', 'MinGW-W64 gcc 9.5.0, HAL 대역 헤더'],
        ['FPGA 보드', 'Combo II-DLD, xc7s75fgga484-1'],
        ['FPGA 도구', 'Vivado 2025.2, Icarus Verilog 12.0, fpga-lab-template v2.0.2'],
    ], [0.92, 2.30], size=7.0)
    add_text(doc, 'MCU 프로젝트는 수업에서 지정한 태그 v1.1.0을 받아 CMake의 Debug 구성으로 빌드했다[1]. 매뉴얼은 VS Code의 '
                  'STM32CubeIDE 확장을 쓰지만, 이 PC에는 아직 확장이 없어 같은 CMake 프리셋을 명령줄에서 실행했다. '
                  '보드에 내려받는 과정(F5)은 실험실에서 수행한다.')
    add_text(doc, '보드가 없는 상태에서 동작을 확인하려고 호스트 시뮬레이션을 만들었다. 보드에 올릴 ece2_474_routine.c를 고치지 않고 '
                  'PC용 gcc로 컴파일하되, stm32g4xx_hal.h 자리에 GPIO와 HAL_GetTick만 흉내 내는 대역 헤더를 넣는다. '
                  '메인 루프 한 바퀴를 50 µs로 두고 polling_routine()을 반복 호출하면서, 시나리오 표의 시각에 버튼 입력을 바꾸고 '
                  '출력 핀이 바뀐 시각을 기록한다. 실제 보드의 루프는 이보다 빠르지만 판단은 ms 단위 tick으로 하므로 결과 시각은 같다.')
    add_text(doc, 'FPGA 실습은 매뉴얼에 실린 설계 모듈과 테스트벤치를 fpga-lab-template 폴더 구조에 넣어 Icarus Verilog로 실행하고, '
                  '같은 파일을 Vivado 2025.2 시뮬레이터에서 다시 실행해 파형 창을 캡처했다. 캡처는 한 단 폭에서 값이 읽히도록 '
                  '시간축을 반으로 잘라 두 줄로 배치했다. 두 설계 모두 조합회로라 클록이 없다.')

    # ---- 3 출처
    add_heading(doc, '3 소스 출처와 실행 스냅샷', 1)
    add_table(doc, ['항목', '식별 정보'], [
        ['MCU 템플릿', f'{UPSTREAM}.git'],
        ['태그·커밋', f'v1.1.0 · {UPSTREAM_COMMIT}'],
        ['FPGA 템플릿', 'Glaysia/fpga-lab-template v2.0.2'],
        ['제출 저장소', REPO],
        ['LAB4 코드', f'{CODE_COMMIT} (보고서의 GitHub 주소가 가리키는 위치)'],
        ['폴더', 'lab4/stm32_lab, stm32_0_hostsim, lab4_01_mult_4bit, lab4_02_bcd_conv'],
    ], [0.9, 2.32], size=7.0)
    add_link_line(doc, 'GitHub', gh(lab4, tree=True))
    add_text(doc, '템플릿에서 바꾼 파일은 Core/Src/ece2_474_routine.c 하나다. 바뀐 내용 전체는 아래 diff에 있고, '
                  '호스트 시뮬레이션의 원자료(CSV)와 FPGA 시뮬레이션 로그도 같은 저장소에 있다. '
                  '시연 영상은 실험 후 보고서에 링크한다.')
    add_link_line(doc, 'diff', gh(f'{lab4}/report/pre/logs/stm32_routine.diff'))

    # ================================================================= STM32
    add_heading(doc, 'STM32_0 폴링 구조와 제공 예제', 1)
    add_heading(doc, '설계 목표와 실행 흐름', 2)
    add_text(doc, '목표는 버튼을 읽고 LED의 동작을 프로그램으로 바꾸는 것이다. 인터럽트는 쓰지 않는다. main.c의 while(1)이 '
                  'polling_routine()을 쉬지 않고 호출하고, 함수는 호출될 때마다 현재 버튼 값과 현재 시각을 읽어 할 일이 있는지 판단한 뒤 '
                  '바로 돌아온다. 기다리는 함수(HAL_Delay)를 쓰지 않으므로 LED 간격을 재는 동안에도 버튼을 놓치지 않는다.')
    add_text(doc, '시각은 HAL_GetTick()으로 읽는다. HAL은 SysTick 인터럽트에서 1 ms마다 tick 변수를 1씩 올리고 HAL_GetTick()은 '
                  '그 값을 돌려준다[4]. SysTick은 Cortex-M 프로세서에 들어 있는 24비트 감산 타이머로, 이런 주기적 시간 기준을 만드는 용도로 '
                  '쓴다[2]. 함수 안의 static 변수는 호출이 끝나도 값이 남으므로 이전 버튼 상태와 마지막 변경 시각을 여기에 둔다.')
    add_table(doc, ['핀', '이름', '역할'], [
        ['PC13', 'BUTTON_ece2', '사용자 버튼 입력. 누르면 HIGH'],
        ['PA5', 'LED_ece2', '내장 LED. HIGH에서 켜짐'],
        ['PB0', 'GPIO_OUT1_ece2', '범용 출력. 과제2의 외부 LED'],
        ['PB1', 'GPIO_OUT2_ece2', '범용 출력. PB0의 대안'],
    ], [0.5, 1.12, 1.6], size=7.2)
    add_text(doc, 'PC13은 CubeMX 설정에서 내부 풀업·풀다운 없이 상승 에지 인터럽트 입력으로 잡혀 있고, 제공 예제도 HIGH를 누름으로 본다. '
                  '폴링 실습에서는 인터럽트 함수(button_isr_p7)를 비워 두고 HAL_GPIO_ReadPin으로 값만 읽는다.')

    add_heading(doc, '제공 예제 1·2의 동작', 2)
    add_table(doc, ['예제', '판단 조건', '동작'], [
        ['1', 'now − last ≥ 1000', 'led_state를 반전해 WritePin. 주기 2초'],
        ['2', '이전 LOW, 현재 HIGH, 직전 인정 후 200 ms 이상', 'TogglePin 한 번'],
    ], [0.42, 1.5, 1.3], size=7.2)
    add_text(doc, '예제 1은 시간이 조건이고 예제 2는 입력의 변화가 조건이다. 예제 2에서 현재 값만 보면 누르고 있는 동안 계속 토글되므로, '
                  '이전 값과 비교해 LOW에서 HIGH로 바뀐 한 순간만 고른다. 뺄셈 now − last는 부호 없는 32비트 연산이라 '
                  'tick이 0xFFFFFFFF에서 0으로 넘어가도 경과 시간이 그대로 나온다.')
    fig = add_figure(doc, IMG / 'hostsim' / 'examples.png',
                     '제공 예제 1·2의 호스트 시뮬레이션. 위는 1초마다 반전, 아래는 누른 순간에만 반전(1.5초 누름도 한 번).', fig)

    # ---- 과제1
    add_heading(doc, 'STM32_0 과제1 버튼으로 LED 간격 변경', 1)
    add_heading(doc, '요구 조건과 설계', 2)
    bullets(doc, ['버튼을 누를 때마다 LED 상태 변경 간격이 1 → 2 → 4 → 8 → 1초로 바뀐다.',
                  '버튼을 누르지 않아도 선택한 간격으로 계속 깜빡인다.',
                  '버튼을 계속 누르고 있어도 간격은 한 번만 바뀐다.'])
    add_text(doc, '매뉴얼의 안내대로 두 제공 예제를 합치고 간격을 저장하는 변수 interval_ms 하나를 더했다. 예제 2의 누름 검출이 '
                  'LED를 토글하던 자리에서 interval_ms를 바꾸고, 예제 1의 고정값 1000이 있던 자리에 interval_ms를 넣는다. '
                  '두 판단은 서로 기다리지 않고 매 호출마다 차례로 실행된다.')
    fig = add_figure(doc, IMG / 'block' / 'stm32_polling_block.png',
                     '과제1의 처리 흐름. 버튼 쪽은 간격만 바꾸고, 시간 쪽은 그 간격으로 LED를 토글한다.', fig)
    add_table(doc, ['static 변수', '초기값', '역할'], [
        ['previous', 'RESET', '직전 호출에서 읽은 버튼 값'],
        ['last_press_ms', '0', '마지막으로 누름을 인정한 시각'],
        ['last_toggle_ms', '0', '마지막으로 LED를 바꾼 시각'],
        ['interval_ms', '1000', '현재 간격. 1000·2000·4000·8000'],
    ], [0.95, 0.55, 1.72], size=7.2)
    add_text(doc, '간격은 2배씩 늘다가 8000 이상이면 1000으로 돌아간다. 값 네 개를 배열에 두고 번호를 돌리는 방법도 있지만, '
                  '수열이 2의 거듭제곱이라 곱셈 한 번이 더 짧고 변수도 하나로 끝난다.')
    add_heading(doc, '핵심 코드', 2)
    add_code(doc, TASK1_CODE, size=6.4, link=gh(routine, 69, 93))
    add_text(doc, 'LED는 WritePin 대신 TogglePin으로 바꾼다. 현재 상태를 따로 기억할 필요가 없고, 포트와 핀 이름을 BLINK_LED_* 한 곳에서 '
                  '정하므로 과제2에서 이 함수는 손대지 않는다.')
    add_text(doc, '간격을 바꿔도 last_toggle_ms는 그대로 둔다. 따라서 바꾼 직후의 첫 변화는 누른 시각이 아니라 직전 LED 변화 시각에서 '
                  '새 간격만큼 뒤에 온다. 8초에서 1초로 돌아갈 때 이미 1초가 지났다면 다음 호출에서 바로 바뀐다.')

    add_heading(doc, '예상 결과: 간격 순환', 2)
    add_text(doc, '버튼을 4.5, 10.5, 22.5, 46.5초에 150 ms씩 누르는 시나리오를 실행했다. LED가 바뀐 시각은 다음과 같다.')
    add_table(doc, ['구간', 'LED가 바뀐 시각 (s)', '간격'], [
        ['시작~4.5 s', '1, 2, 3, 4', '1 s'],
        ['4.5~10.5 s', '6, 8, 10', '2 s'],
        ['10.5~22.5 s', '14, 18, 22', '4 s'],
        ['22.5~46.5 s', '30, 38, 46', '8 s'],
        ['46.5 s 이후', '47, 48, 49, 50, 51, 52', '1 s'],
    ], [0.92, 1.6, 0.7], size=7.2)
    fig = add_figure(doc, IMG / 'hostsim' / 'task1_cycle.png',
                     '과제1 호스트 시뮬레이션. 누를 때마다 LED 변화 간격이 2, 4, 8초로 늘고 네 번째 누름에서 1초로 돌아간다.', fig)
    add_link_line(doc, '원자료', gh(f'{lab4}/stm32_0_hostsim/results/task1_cycle.csv'))
    add_text(doc, '4.5초에 눌렀을 때 다음 변화가 5초가 아니라 6초에 온 것은 위에서 설명한 대로 직전 변화(4초)를 기준으로 2초를 세기 때문이다. '
                  '같은 입력을 tick이 0xFFFFB1E0에서 시작해 20초 뒤 0으로 넘어가도록 바꿔 실행해도 변화 간격은 같았다. '
                  '다만 이 경우 last_toggle_ms의 초기값 0과의 차이가 커서 시작하자마자 한 번 바뀐다.')

    add_heading(doc, '예상 결과: 길게 누름', 2)
    add_text(doc, '2.5초부터 6초 동안 누르고 있어도 간격은 2초로 한 번만 바뀌고, 누르고 있는 동안에도 LED는 2초마다 계속 바뀐다. '
                  '함수가 버튼을 기다리지 않기 때문이다.')
    fig = add_figure(doc, IMG / 'hostsim' / 'task1_hold.png',
                     '6초 동안 누른 경우. 간격은 1초에서 2초로 한 번만 바뀐다.', fig)

    add_heading(doc, '접점 떨림과 대비안', 2)
    add_text(doc, '기계식 버튼은 누르거나 뗄 때 접점이 잠깐 붙었다 떨어지기를 반복할 수 있다. 에지마다 0.4 ms 간격으로 5번 떨리는 모델을 넣어 '
                  '실행했다. 150 ms 누름에서는 누를 때의 떨림도, 150 ms 뒤 뗄 때의 떨림도 200 ms 잠금 안에 있어 한 번만 인정됐다. '
                  '1.5초 누름에서는 달랐다. 뗄 때는 잠금이 이미 풀려 있어서, 떨림 중 다시 HIGH가 되는 순간이 새 누름으로 인정됐다.')
    fig = add_figure(doc, IMG / 'hostsim' / 'task1_bounce_zoom.png',
                     '떼는 순간(8.000 s)의 떨림 모델. LOW가 된 뒤 다시 HIGH가 되는 상승 에지가 다섯 번 생긴다.', fig)
    fig = add_figure(doc, IMG / 'hostsim' / 'task1_bounce.png',
                     '떨림이 있을 때. 과제1 코드는 두 번째 누름을 뗄 때 간격이 4초에서 8초로 한 번 더 바뀌고, 대비안은 4초를 유지한다.', fig)
    add_text(doc, '대비안은 누름 인정 시각 대신 버튼 값이 마지막으로 바뀐 시각을 기억하고, 그로부터 50 ms 이상 조용했던 뒤의 상승만 '
                  '누름으로 본다. 뗄 때의 하강도 시각을 갱신하므로 바로 뒤따르는 떨림은 걸러진다.', keep_next=True)
    add_code(doc, GUARD_CODE, size=6.4,
             link=gh(f'{lab4}/stm32_0_hostsim/variants/task1_release_guard.c', 20, 31))
    add_text(doc, '이 떨림은 모델이지 측정이 아니다. NUCLEO 보드의 버튼 회로에 떨림을 줄이는 소자가 있으면 과제1 코드만으로 충분할 수 있다. '
                  '그래서 매뉴얼의 구성에 가까운 과제1 코드를 먼저 올리고, 길게 눌렀다 뗄 때 간격이 두 단계 넘어가는 경우에만 '
                  '대비안으로 바꾼다. 대비안은 떨림이 없는 순환·길게 누름 시나리오에서 과제1 코드와 같은 결과를 냈다.')

    add_heading(doc, '빌드 결과', 2)
    add_code(doc, '\n'.join(['$ cmake --preset Debug', '$ cmake --build --preset Debug',
                             'RAM:    2656 B / 128 KB  (2.03 %)', 'FLASH: 25848 B / 512 KB  (4.93 %)',
                             'polling_routine  0x080056d4  152 B']), size=6.4,
             link=gh(f'{lab4}/report/pre/logs/stm32_arm_build.txt'), label='로그')
    add_text(doc, 'Arm GNU Toolchain 13.3으로 빌드했고 ece2_474_routine.c에서 경고는 없었다. interval_ms는 0이 아닌 초기값이 있어 .data에, 나머지 세 변수는 .bss에 놓였다.')

    add_heading(doc, '보드에서 확인할 항목', 2)
    add_table(doc, ['확인 항목', '예상 관찰'], [
        ['제공 예제 1', '내장 LED가 1초 켜지고 1초 꺼짐'],
        ['제공 예제 2', '누를 때마다 LED 상태가 한 번 바뀜'],
        ['과제1 시작', '리셋 후 1초 간격'],
        ['누름 1~3회', '2초, 4초, 8초 간격. 스톱워치로 확인'],
        ['누름 4회', '다시 1초 간격'],
        ['길게 누름', '간격이 한 단계만 바뀜. 두 단계면 대비안 적용'],
        ['누르는 동안', 'LED가 멈추지 않고 계속 깜빡임'],
        ['다운로드', 'Build만으로는 바뀌지 않음. 디버깅 종료 후 F5'],
    ], [0.9, 2.32], size=7.2)

    # ---- 과제2
    add_heading(doc, 'STM32_0 과제2 외부 LED로 동작시키기', 1)
    add_heading(doc, '회로 구성', 2)
    add_text(doc, '과제1의 동작을 PB0에 연결한 외부 LED로 옮긴다. PB0는 push-pull 출력으로 설정되어 있고 초기값은 LOW다. '
                  '핀에서 저항과 LED를 거쳐 GND로 가는 직렬 회로이므로 PB0가 HIGH일 때 켜진다. 내장 LED와 극성이 같아 코드의 논리를 뒤집을 필요가 없다.')
    fig = add_figure(doc, IMG / 'block' / 'stm32_led_wiring.png',
                     '과제2 결선. PB0 → 저항 → LED 긴 다리, LED 짧은 다리 → GND. 커넥터 위치는 보드 설명서로 다시 확인한다.', fig)
    add_table(doc, ['저항', '계산', 'LED 전류'], [
        ['330 Ω', '(3.3 − 2.0) / 330', '약 3.9 mA'],
        ['1 kΩ', '(3.3 − 2.0) / 1000', '약 1.3 mA'],
    ], [0.7, 1.52, 1.0], size=7.2)
    add_text(doc, '적색 LED의 순방향 전압을 약 2.0 V로 가정한 값이다. 저항이 클수록 어둡지만 두 값 모두 켜진 것을 알아보기에 충분하다. '
                  '저항 없이 연결하면 전류를 제한하는 것이 없으므로 반드시 저항을 넣는다.')
    add_heading(doc, '코드 변경', 2)
    add_text(doc, 'polling_routine()은 그대로 두고, 파일 위쪽에서 LED 핀을 고르는 매크로 값만 0에서 1로 바꾼다.', keep_next=True)
    add_code(doc, TASK2_CODE, size=6.4, link=gh(routine, 54, 64))
    add_text(doc, 'main.h에서 PB0의 이름은 GPIO_OUT1_ece2_Pin과 GPIO_OUT1_ece2_GPIO_Port다. PB1을 쓰려면 GPIO_OUT2_ece2로 바꾼다. '
                  '값을 1로 두고 Arm GCC로 컴파일한 결과에서 HAL_GPIO_TogglePin의 핀 인자가 0x20(5번 핀)에서 0x01(0번 핀)로 바뀐 것을 확인했다.')
    add_heading(doc, '예상 결과', 2)
    add_text(doc, '과제1과 같은 시나리오에서 PB0가 바뀐 시각은 과제1의 PA5와 모두 같았고, PA5는 LOW로 남았다.')
    fig = add_figure(doc, IMG / 'hostsim' / 'task2_cycle.png',
                     '과제2 호스트 시뮬레이션. 외부 LED 핀 PB0가 과제1과 같은 시각에 바뀐다.', fig)
    add_heading(doc, '보드에서 확인할 항목', 2)
    add_table(doc, ['확인 항목', '예상 관찰'], [
        ['결선', 'PB0 → 저항 → LED(+), LED(−) → GND'],
        ['과제1 먼저', '내장 LED로 과제1을 확인한 뒤 진행'],
        ['외부 LED', '리셋 후 1초 간격으로 깜빡임'],
        ['버튼', '외부 LED도 1→2→4→8→1초로 바뀜'],
        ['내장 LED', '꺼진 채로 유지'],
        ['안 켜질 때', 'LED 방향, GND 연결, 매크로 값, F5 다운로드 순으로 확인'],
    ], [0.9, 2.32], size=7.2)

    # ================================================================= FPGA 01
    add_heading(doc, 'FPGA-01 4비트 곱셈기', 1)
    add_heading(doc, '회로 목적과 구조', 2)
    add_text(doc, '부호 없는 4비트 수 a와 b를 곱해 8비트 m을 낸다. 가장 큰 결과가 15 × 15 = 225이므로 8비트(최대 255)면 넘치지 않는다. '
                  '입력만으로 출력이 정해지는 조합회로이며 클록과 리셋이 없다.')
    add_text(doc, '이진 곱셈은 손으로 하는 곱셈과 같다. 승수의 각 비트와 피승수를 AND해 부분곱을 만들고, 한 자리씩 옮겨 더한다. '
                  '4비트끼리의 곱은 AND 게이트 16개와 4비트 가산기 3개로 구성할 수 있다[3]. 설계에서는 이를 직접 그리지 않고 '
                  'a * b로 적어 합성 도구가 회로를 만들게 했다.')
    fig = add_figure(doc, IMG / 'block' / '01_mult_block.png', 'FPGA-01 블록 흐름. 설계 top은 mult_4bit.', fig)
    add_code(doc, MULT_CODE, size=6.4, link=gh(f'{lab4}/lab4_01_mult_4bit/src', tree=True))
    add_text(doc, '곱셈은 multiply_unsigned가 하고 mult_4bit은 폭을 4로 정해 연결만 한다. 출력 폭을 2 × WIDTH로 선언했기 때문에 '
                  '곱의 상위 비트가 잘리지 않는다.')
    add_heading(doc, '파일의 역할', 2)
    d1 = f'{lab4}/lab4_01_mult_4bit'
    files1 = [['src/multiply_unsigned.v', '폭을 매개변수로 받는 곱셈'],
              ['src/mult_4bit.v', '설계 top. 4비트 입출력 연결'],
              ['sim/tb.sv', '매뉴얼의 TB. 256쌍 전수 검사'],
              ['sim/tb_view.sv', '파형 확인용. 경계값 12쌍을 10 ns씩'],
              ['constraints/mult_4bit.xdc', 'a·b 스위치 8핀, m LED 8핀, LVCMOS33']]
    add_table(doc, ['파일', '역할'], files1, [1.28, 1.94], size=7.0, center_first=False,
              links=[gh(f'{d1}/{f[0]}') for f in files1])
    add_text(doc, 'tb.sv는 매뉴얼의 코드에 파형 기록($dumpfile, $dumpvars) 한 줄만 더했다. 전수 검사는 1 ns마다 입력이 바뀌어 값을 읽기 '
                  '어려우므로, 값을 눈으로 확인할 tb_view.sv를 따로 작성했다.')
    add_heading(doc, '시뮬레이션 결과', 2)
    add_code(doc, '\n'.join(['[Icarus] PASS multiplier: 256 input pairs',
                             '         $finish called at 256000 (1ps)',
                             '[Vivado] PASS multiplier: 256 input pairs',
                             '         $finish called at time : 256 ns',
                             '[Vivado] PASS multiplier view: 12 directed pairs']), size=6.4,
             link=gh(f'{lab4}/report/pre/logs', tree=True), label='로그')
    fig = add_figure(doc, IMG / 'vivado' / '01_mult_view.png',
                     'FPGA-01 Vivado 파형(tb_view, 10진 표시). 캡처를 0~60 ns와 60~120 ns 두 줄로 나눠 놓았다. '
                     '0×15=0, 3×5=5×3=15, 12×10=120, 15×15=225.', fig)
    add_text(doc, '두 시뮬레이터 모두 256쌍을 통과했다. 파형에서 한쪽 입력이 0이면 결과가 0이고, 3×5와 5×3이 같으며, '
                  '최댓값 225에서 m[7:0]이 1110_0001이 되는 것을 확인했다.')
    add_heading(doc, '구현 결과', 2)
    add_table(doc, ['항목', '값'], [
        ['LUT', '16개 (DSP 0, 레지스터 0)'],
        ['가장 긴 경로', 'a[0] → m[7], 12.063 ns'],
        ['DRC', '오류 0'],
        ['bitstream', 'bitstream/mult_4bit.bit'],
    ], [1.0, 2.22], size=7.2)
    add_text(doc, '가장 긴 경로는 최하위 입력에서 최상위 출력까지이며, 부분곱의 자리올림이 끝까지 전달되는 길과 일치한다. '
                  '이 지연은 입출력 버퍼를 포함한 값이고, 스위치를 손으로 바꾸는 실험에서는 의미 있는 제약이 아니다.')
    add_heading(doc, '보드에서 확인할 항목', 2)
    add_table(doc, ['a', 'b', 'm (10진)', 'LED m[7:0]'], [
        ['0000', '1111', '0', '0000_0000'],
        ['0011', '0101', '15', '0000_1111'],
        ['1000', '1000', '64', '0100_0000'],
        ['1100', '1010', '120', '0111_1000'],
        ['1111', '1111', '225', '1110_0001'],
    ], [0.62, 0.62, 0.78, 1.2], size=7.2)
    add_text(doc, 'XDC의 핀은 LAB1의 4비트 가산기(스위치)와 LAB3의 LED 배치를 그대로 썼다. 교재의 곱셈기 핀 표와 다르면 교재를 따른다. '
                  '스위치와 LED의 극성이 반대로 보이면 보드 회로를 확인한다.')

    # ================================================================= FPGA 02
    add_heading(doc, 'FPGA-02 BCD→이진수 변환', 1)
    add_heading(doc, '회로 목적과 구조', 2)
    add_text(doc, 'BCD는 10진수의 각 자리를 4비트로 따로 적는 코드다[3]. 십의 자리 ten과 일의 자리 one을 받아 그 값 ten × 10 + one을 '
                  '7비트 이진수 bin으로 낸다. 유효한 입력의 최댓값이 99(110_0011)이므로 7비트면 충분하다.')
    add_text(doc, '4비트 입력은 10~15도 나타낼 수 있지만 BCD에서는 쓰지 않는 값이다. 매뉴얼의 설계는 두 자리가 모두 0~9일 때만 '
                  'valid를 1로 내도록 원본 예제에 출력을 하나 더했다[1].')
    fig = add_figure(doc, IMG / 'block' / '02_bcd_block.png', 'FPGA-02 블록 흐름. 설계 top은 bcd_conv.', fig)
    add_code(doc, BCD_CODE, size=6.4, link=gh(f'{lab4}/lab4_02_bcd_conv/src/bcd_to_binary.v'))
    add_text(doc, "곱하는 수를 7'd10으로 적어 곱셈이 7비트 폭에서 계산되게 했다. ten × 10은 ten × 8 + ten × 2, 곧 자리 옮김 두 번과 "
                  '덧셈으로 만들 수 있어 실제 곱셈기가 필요하지 않다. top인 bcd_conv는 포트 순서(one, ten)만 바꿔 연결한다.')
    add_heading(doc, '범위를 벗어난 입력', 2)
    add_table(doc, ['ten', 'one', '계산값', 'bin', 'valid'], [
        ['9', '9', '99', '99', '1'],
        ['10', '0', '100', '100', '0'],
        ['0', '10', '10', '10', '0'],
        ['12', '7', '127', '127', '0'],
        ['15', '15', '165', '37', '0'],
    ], [0.5, 0.5, 0.74, 0.74, 0.74], size=7.2)
    add_text(doc, 'valid가 0일 때도 bin에는 계산값이 그대로 나오고, 127을 넘으면 하위 7비트만 남는다(165 − 128 = 37). '
                  '따라서 bin은 valid가 1일 때만 믿을 수 있다. 일의 자리에 10을 넣으면 bin이 10이 되어 유효한 입력 "1, 0"과 구별되지 않는데, '
                  '이를 가르는 것이 valid다.')
    add_heading(doc, '파일의 역할', 2)
    d2 = f'{lab4}/lab4_02_bcd_conv'
    files2 = [['src/bcd_to_binary.v', '변환과 유효성 검사'],
              ['src/bcd_conv.v', '설계 top. 입출력 연결'],
              ['sim/tb.sv', '매뉴얼의 TB. 유효 100개, 무효 156개'],
              ['sim/tb_view.sv', '파형 확인용. 유효 8개, 무효 4개'],
              ['constraints/bcd_conv.xdc', 'ten·one 스위치 8핀, bin·valid LED 8핀']]
    add_table(doc, ['파일', '역할'], files2, [1.28, 1.94], size=7.0, center_first=False,
              links=[gh(f'{d2}/{f[0]}') for f in files2])
    add_heading(doc, '시뮬레이션 결과', 2)
    add_code(doc, '\n'.join(['[Icarus] PASS BCD: 100 valid and 156 invalid inputs',
                             '         $finish called at 256000 (1ps)',
                             '[Vivado] PASS BCD: 100 valid and 156 invalid inputs',
                             '         $finish called at time : 256 ns',
                             '[Vivado] PASS BCD view: 12 directed inputs']), size=6.4,
             link=gh(f'{lab4}/report/pre/logs', tree=True), label='로그')
    fig = add_figure(doc, IMG / 'vivado' / '02_bcd_view.png',
                     'FPGA-02 Vivado 파형(tb_view, 10진 표시, 두 줄 배치). 80 ns까지는 유효 입력으로 valid=1, 이후 네 입력은 valid=0.', fig)
    fig = add_figure(doc, IMG / 'vivado' / '02_bcd_full.png',
                     'FPGA-02 Vivado 전체 파형(tb, 256 ns, 두 줄 배치). ten이 0~9인 동안 one이 0~9인 구간에서만 valid가 1이고, '
                     'ten이 10 이상이면 계속 0이다. one과 bin은 1 ns마다 바뀌어 띠로 보인다.', fig)
    add_text(doc, '매뉴얼의 TB는 valid를 256개 입력 모두에서, bin을 유효한 100개에서 검사한다. 무효 입력의 bin은 검사하지 않으므로 '
                  'tb_view.sv에서 계산값의 하위 7비트와 같은지 따로 확인했다.')
    add_heading(doc, '구현 결과', 2)
    add_table(doc, ['항목', '값'], [
        ['LUT', '7개 (DSP 0, 레지스터 0)'],
        ['가장 긴 경로', 'ten[0] → bin[6], 11.226 ns'],
        ['DRC', '오류 0'],
        ['bitstream', 'bitstream/bcd_conv.bit'],
    ], [1.0, 2.22], size=7.2)
    add_heading(doc, '보드에서 확인할 항목', 2)
    add_table(doc, ['ten', 'one', 'bin[6:0]', 'valid'], [
        ['0000', '0000', '000_0000 (0)', '켜짐'],
        ['0100', '0010', '010_1010 (42)', '켜짐'],
        ['1001', '1001', '110_0011 (99)', '켜짐'],
        ['0100', '1010', '011_0010 (50)', '꺼짐'],
        ['1010', '0000', '110_0100 (100)', '꺼짐'],
    ], [0.62, 0.62, 1.28, 0.7], size=7.2)
    add_text(doc, 'valid는 LED 8개 중 가장 왼쪽(곱셈기의 m[7] 자리)에 두었다. 교재의 핀 표와 다르면 교재를 따른다.')

    # ---- 결론
    add_heading(doc, '4 예상 결과 정리와 실험 계획', 1)
    add_text(doc, 'MCU와 FPGA 실습은 같은 일을 다른 방식으로 한다. 곱셈기와 BCD 변환기는 입력이 바뀌면 회로가 곧바로 출력을 내고, '
                  'STM32에서는 프로그램이 입력과 시각을 반복해서 읽어 출력을 바꾼다. 과제1에서 버튼 처리와 LED 처리가 서로 막지 않는 것은 '
                  '함수가 기다리지 않고 매번 돌아오기 때문이며, 이 구조가 다음 주의 인터럽트 방식과 비교할 기준이 된다.')
    add_table(doc, ['순서', '할 일', '기록할 것'], [
        ['1', '제공 예제 1·2 실행', 'LED 동작 영상'],
        ['2', '과제1 다운로드', '네 간격을 스톱워치로 측정'],
        ['3', '길게 누름 시험', '간격이 몇 단계 바뀌는지'],
        ['4', '과제2 결선·다운로드', '외부 LED 영상, 결선 사진'],
        ['5', '곱셈기 bitstream', '표의 다섯 입력과 LED'],
        ['6', 'BCD 변환 bitstream', '표의 다섯 입력과 LED'],
    ], [0.42, 1.4, 1.4], size=7.2)
    add_text(doc, '실험 후 보고서에서는 위 측정값을 이 보고서의 예상 시각·예상 LED 값과 비교하고, 길게 누름에서 떨림 문제가 실제로 나왔는지와 '
                  '그때의 해결 과정을 적는다. 제출 버전의 태그와 커밋 해시, 시연 영상 링크도 그때 함께 적는다.')

    add_heading(doc, '참고 문헌', 1)
    add_text(doc, '[1] 이해리, 「STM32 실습 매뉴얼」(STM32_0, FPGA 추가 실습) 및 「전자전기컴퓨터설계실험Ⅱ STM32 실습 안내」, '
                  f'서울시립대학교, 2026. {UPSTREAM} (v1.1.0).', size=9.0)
    add_text(doc, '[2] Joseph Yiu, The Definitive Guide to ARM Cortex-M3 and Cortex-M4 Processors, 3rd ed., Newnes, 2013. '
                  'ISBN 978-0-12-408082-9.', size=9.0)
    add_text(doc, '[3] M. Morris Mano and Michael D. Ciletti, Digital Design: With an Introduction to the Verilog HDL, VHDL, '
                  'and SystemVerilog, 6th ed., Pearson, 2018. ISBN 978-0-13-454989-7.', size=9.0)
    add_text(doc, '[4] STMicroelectronics, STM32CubeG4 MCU Firmware Package (stm32g4xx_hal.c, stm32g4xx_hal_gpio.c). '
                  'https://github.com/STMicroelectronics/STM32CubeG4', size=9.0)

    doc.core_properties.title = 'LAB4 예비 보고서'
    doc.core_properties.subject = 'STM32_0 LED·버튼 폴링, FPGA 4비트 곱셈기·BCD→이진수 변환'
    doc.core_properties.author = '김태이'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
